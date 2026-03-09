import openpyxl
from openpyxl import Workbook
import pandas
import sys
import datetime

args = sys.argv

sys.path += ['/opt/iris/lib/python','/opt/iris/mgr/python']

import iris

inputFilename = args[1]
outputFilename = args[2]
reportmonth = args[3]
reportday = args[4]

year = datetime.datetime.today().year
reiwayear = year - 2018

status = iris.cls('YNC.Expense')._KillExtent()
wb = openpyxl.load_workbook(inputFilename)
ws = wb['input']
		
row_index = 0
for row in ws.iter_rows():
	row_index = row_index + 1
	if row_index < 5: continue
	month = ws.cell(row=row_index,column=2).value
	if (month == '' or month is None): continue
	day = ws.cell(row=row_index,column=3).value
	if (day == '' or day is None): continue
	paymentto = ws.cell(row=row_index,column=4).value
	accounts = ws.cell(row=row_index, column=5).value
	if accounts is None: accounts = '旅費交通費'
	amount = ws.cell(row=row_index, column=6).value
	if (amount == '' or amount is None): amount = 0
	description = ws.cell(row=row_index, column=7).value
	if description is None: description = 'no data'
	isjreimbursement = ws.cell(row=row_index,column=8).value
	onbehalf = ws.cell(row=row_index,column=9).value
	if isjreimbursement is None: isjreimbursement = ''
	if onbehalf is None: onbehalf = ''
	sql = iris.sql.prepare("insert into ync.expense(reportmonth,reportday,paymentto,accounts,amount,description,isjreimbursement,onbehalf) values(?,?,?,?,?,?,?,?)")
	sql.execute(month,day,paymentto,accounts,amount,description,isjreimbursement,onbehalf)

	sql = iris.sql.prepare("select description from ync.expenseitem where description = ?")
	rs = sql.execute(description)
	try:
		next(rs)
		exist = True
	except Exception:
		exist = False
	if exist is False:
		sql = iris.sql.prepare("insert into ync.expenseitem(paymentto,accounts,amount,description) values(?,?,?,?)")
		sql.execute(paymentto,accounts,amount,description)
	else:
		if not (amount == 0 or amount is None):
			sql = iris.sql.prepare("update ync.expenseitem set paymentto=?,accounts=?,amount=? where description= ?")
			sql.execute(paymentto,accounts,amount,description)

wb.close()

sql = None

wb = openpyxl.load_workbook(outputFilename)
		
ws = wb['印刷用']
		
itemline1 = iris.sql.exec("select max(reportmd) from ync.expense where accounts = '旅費交通費'").dataframe()

for index,row in itemline1.iterrows():
	rowline = list(row)
	maxmd = rowline[0]
	maxmd = str(maxmd)
	if (len(maxmd) == 3): maxmd = '0' + maxmd
	maxmonth = maxmd[0:2]
	maxday = maxmd[2:4]

minmonth = reportmonth
minday = reportday

maxdt = datetime.datetime(year=year, month=int(maxmonth), day=int(maxday))
mindt = datetime.datetime(year=year, month=int(minmonth), day=int(minday))

td = maxdt - mindt
nights = td.days

ws.cell(row=2,column=2).value = reiwayear
ws.cell(row=2,column=4).value = maxmonth
ws.cell(row=2,column=6).value = maxday

ws.cell(row=7,column=14).value = reiwayear
ws.cell(row=7,column=16).value = minmonth
ws.cell(row=7,column=18).value = minday
ws.cell(row=8,column=14).value = reiwayear
ws.cell(row=8,column=16).value = maxmonth
ws.cell(row=8,column=18).value = maxday

ws.cell(row=13,column=4).value = str(reiwayear) + '年' + minmonth + '月' + minday + '日'
ws.cell(row=13,column=12).value = str(reiwayear) + '年' + maxmonth + '月' + maxday + '日'

sql = iris.sql.prepare("Select reportmonth, reportday, amount, description, paymentto from ync.expense where (reportmonth = ? and reportday >= ? and accounts = '旅費交通費') order by reportmonth, reportday")
itemline = sql.execute(reportmonth,reportday).dataframe()

hotelamount = 0
hotelname = ''

# まず書き込み対象の行リストを組み立てる
# 各要素は (month, day, amount, description) のdict
items = []

for index,row in itemline.iterrows():
	rowline = list(row)
	month = rowline[0]
	day = rowline[1]
	amount = rowline[2]
	description = rowline[3]
	paymentto = rowline[4]

	sql = iris.sql.prepare("select accounts, amount from ync.expenseitem where description = ?")
	rs = sql.execute(description)
	for index,row in enumerate(rs):
		accounts = row[0]
		if (accounts is None or accounts == ''): accounts = '旅費交通費'
		if (amount == 0 or amount is None): amount = row[1]

	if description == '雲雀丘花屋敷>大阪梅田': continue
	if description == '大阪梅田>雲雀丘花屋敷': continue

	if description == 'HOTEL':
		hotelamount = hotelamount + amount
		hotelname = hotelname + paymentto + ' '
		continue

	if paymentto == 'JAL':
		transport_type = '飛行機'
	elif paymentto == 'PEACH':
		transport_type = '飛行機'
	elif paymentto == 'SKYMARK':
		transport_type = '飛行機'
	elif paymentto == 'ANA':
		transport_type = '飛行機'
	else:
		transport_type = '電車'

	full_description = transport_type + '　(' + description + ')'
	items.append({'month': month, 'day': day, 'amount': amount, 'description': full_description, 'transport_type': transport_type, 'route': description})


def merge_same_day(items, merge_count):
	"""同じ日付かつ同じ交通手段の行を merge_count 件ずつ1行にまとめる処理を1パス行う"""
	merged = []
	i = 0
	while i < len(items):
		base = items[i]
		# 同じ日付・同じ交通手段の連続する行をmerge_count件集める
		group = [base]
		j = i + 1
		while j < len(items) and len(group) < merge_count:
			if (items[j]['month'] == base['month']
				and items[j]['day'] == base['day']
				and items[j]['transport_type'] == base['transport_type']):
				group.append(items[j])
				j += 1
			else:
				break
		if len(group) >= 2:
			# ルート部分を結合する
			# 隣接する区間で到着駅と次の出発駅が同じ場合は繋げて A>B>C 形式にする
			def connect_routes(routes):
				result = routes[0]
				for r in routes[1:]:
					prev_arrival = result.split('>')[-1]
					next_departure = r.split('>')[0]
					if prev_arrival == next_departure:
						result = result + '>' + '>'.join(r.split('>')[1:])
					else:
						result = result + ' ' + r
				return result
			merged_route = connect_routes([item['route'] for item in group])
			merged_desc = base['transport_type'] + '　(' + merged_route + ')'
			merged_amount = sum(item['amount'] for item in group)
			merged.append({
				'month': base['month'],
				'day': base['day'],
				'amount': merged_amount,
				'description': merged_desc,
				'transport_type': base['transport_type'],
				'route': merged_route
			})
			i = j
		else:
			merged.append(base)
			i += 1
	return merged


MAX_LINES = 10

# 10行を超える場合、同じ日付の2行→3行→...とマージを繰り返す
merge_count = 2
while len(items) > MAX_LINES:
	prev_len = len(items)
	items = merge_same_day(items, merge_count)
	# 同じmerge_countで変化がなくなったら merge_count を増やす
	if len(items) == prev_len:
		merge_count += 1
		if merge_count > 100:
			# 無限ループ防止: これ以上マージできない場合は打ち切り
			print('# of lineitems exceeded the max limit even after merging')
			break

# 行リストをシートに書き込む
linepos = 16
for item in items:
	linepos += 1
	ws.cell(row=linepos, column=1).value = str(item['month']) + '/' + str(item['day'])
	ws.cell(row=linepos, column=6).value = str(item['month']) + '/' + str(item['day'])
	ws.cell(row=linepos, column=10).value = item['description']
	ws.cell(row=linepos, column=16).value = item['amount']

sql = None
ws.cell(row=30,column=8).value = hotelamount
ws.cell(row=30,column=2).value = hotelname
ws.cell(row=28,column=5).value = nights
wb.save(outputFilename)
wb.close()
