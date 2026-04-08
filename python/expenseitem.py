import openpyxl
import sys

sys.path += ['/opt/iris/lib/python', '/opt/iris/mgr/python']
import iris

def do_import(excel_file):
    # 全インスタンス削除
    iris.cls('YNC.ExpenseItem')._KillExtent()

    wb = openpyxl.load_workbook(excel_file)
    ws = wb.active

    # 先頭行からカラム名を取得
    headers = []
    for cell in ws[1]:
        headers.append(cell.value)

    # IDカラムのインデックスを特定（IDはIRISが自動採番するため挿入対象外）
    insert_cols = [(i, h) for i, h in enumerate(headers) if h and h.upper() != 'ID']

    # INSERT文をカラム名から動的に生成
    col_names = ', '.join(h for _, h in insert_cols)
    placeholders = ', '.join('?' for _ in insert_cols)
    sql = iris.sql.prepare(
        f"INSERT INTO YNC.ExpenseItem ({col_names}) VALUES ({placeholders})"
    )

    for row in ws.iter_rows(min_row=2, values_only=True):
        # 全カラムがNoneの行はスキップ
        if all(v is None for v in row):
            continue
        values = [row[i] if row[i] is not None else '' for i, _ in insert_cols]
        sql.execute(*values)

    wb.close()
    print(f"Import completed: {excel_file}")


def do_export(excel_file):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'download'

    # エクスポート対象ファイルが既に存在する場合、ヘッダからカラム順を取得
    try:
        wb_existing = openpyxl.load_workbook(excel_file)
        ws_existing = wb_existing.active
        headers = [cell.value for cell in ws_existing[1] if cell.value]
        wb_existing.close()
    except FileNotFoundError:
        # ファイルがない場合はデフォルト順
        headers = ['ID', 'Accounts', 'Amount', 'Description', 'OnBeHalf', 'PaymentTo']

    ws.append(headers)

    # SELECTのカラム順をヘッダに合わせる
    # IDはIRISの%IDとして取得
    select_cols = []
    for h in headers:
        if h.upper() == 'ID':
            select_cols.append('%ID AS ID')
        else:
            select_cols.append(h)

    col_list = ', '.join(select_cols)
    rs = iris.sql.exec(f"SELECT {col_list} FROM YNC.ExpenseItem ORDER BY Description")

    for row in rs:
        ws.append(list(row))

    wb.save(excel_file)
    wb.close()
    print(f"Export completed: {excel_file}")


if __name__ == '__main__':
    if len(sys.argv) < 3:
        print("Usage: python3 expenseitem.py import|export <excel_file>")
        sys.exit(1)

    command = sys.argv[1].lower()
    excel_file = sys.argv[2]

    if command == 'import':
        do_import(excel_file)
    elif command == 'export':
        do_export(excel_file)
    else:
        print(f"Unknown command: {command}. Use 'import' or 'export'.")
        sys.exit(1)
