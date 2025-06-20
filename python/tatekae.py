import openpyxl
from openpyxl import Workbook
import pandas
import sys
import datetime
import os

args = sys.argv

sys.path += ['/opt/iris/lib/python','/opt/iris/mgr/python']

dir = args[1]

files = os.listdir(dir + '/')
nittous = []
for file in files:
	if '経費' in file:   
		wb = openpyxl.load_workbook(file, data_only=True)
		ws = wb['sorted']
		keihitotal = ws.cell(row=100,column=10).value
		keihitotal2 = ws['J100'].value
		
		wb.close()
	if '出張報告精算書' in file:
		wb = openpyxl.load_workbook(file, data_only=True)
		ws = wb['印刷用']
		price = ws.cell(row=28,column=8).value
		month = ws.cell(row=7,column=16).value
		date = ws.cell(row=7,column=18).value
		nittou = [month,date,price]
		nittous.append(nittou)
		wb.close()
count = 0
outputfile = dir + '/立て替え.xlsx'
wb = openpyxl.load_workbook(outputfile)
ws = wb['Sheet1']
for nittou in nittous:
	count = count + 1
	ws.cell(row=count,column=1).value = nittou[0]
	ws.cell(row=count,column=2).value = nittou[1]
	ws.cell(row=count,column=3).value = nittou[2]
ws.cell(row=19,column=3).value = keihitotal
wb.save(outputfile)
wb.close()
