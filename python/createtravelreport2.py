import openpyxl
from openpyxl import Workbook
import pandas
import sys
import datetime

args = sys.argv

sys.path += ['/opt/iris/lib/python','/opt/iris/mgr/python']

import iris

inputFilename = args[1]
outputFilename = args[2]
reportmonth = args[3]
reportday = args[4]

year = datetime.datetime.today().year
reiwayear = year - 2018

status = iris.cls('YNC.Expense')._KillExtent()
wb = openpyxl.load_workbook(inputFilename)
ws = wb['input']
		
row_index = 0
for row in ws.iter_rows():
	row_index = row_index + 1
	if row_index < 5: continue
	month = ws.cell(row=row_index,column=2).value
	if (month == '' or month is None): continue
	day = ws.cell(row=row_index,column=3).value
	if (day == '' or day is None): continue
	paymentto = ws.cell(row=row_index,column=4).value
	accounts = ws.cell(row=row_index, column=5).value
	if accounts is None: accounts = '旅費交通費'
	amount = ws.cell(row=row_index, column=6).value
	if (amount == '' or amount is None): amount = 0
	description = ws.cell(row=row_index, column=7).value
	if description is None: description = 'no data'
	isjreimbursement = ws.cell(row=row_index,column=8).value
	onbehalf = ws.cell(row=row_index,column=9).value
	if isjreimbursement is None: isjreimbursement = ''
	if onbehalf is None: onbehalf = ''
	sql = iris.sql.prepare("insert into ync.expense(reportmonth,reportday,paymentto,accounts,amount,description,isjreimbursement,onbehalf) values(?,?,?,?,?,?,?,?)")
	sql.execute(month,day,paymentto,accounts,amount,description,isjreimbursement,onbehalf)

	sql = iris.sql.prepare("select description from ync.expenseitem where description = ?")
	rs = sql.execute(description)
	try:
		next(rs)
		exist = True
	except Exception:
		exist = False
	if exist is False:
		sql = iris.sql.prepare("insert into ync.expenseitem(paymentto,accounts,amount,description) values(?,?,?,?)")
		sql.execute(paymentto,accounts,amount,description)
	else:
		if not (amount == 0 or amount is None):
			sql = iris.sql.prepare("update ync.expenseitem set paymentto=?,accounts=?,amount=? where description= ?")
			sql.execute(paymentto,accounts,amount,description)

wb.close()

sql = None

wb = openpyxl.load_workbook(outputFilename)
		
ws = wb['印刷用']
		
itemline1 = iris.sql.exec("select max(reportmd) from ync.expense where accounts = '旅費交通費'").dataframe()

for index,row in itemline1.iterrows():
	rowline = list(row)
	maxmd = rowline[0]
	maxmd = str(maxmd)
	if (len(maxmd) == 3): maxmd = '0' + maxmd
	maxmonth = maxmd[0:2]
	maxday = maxmd[2:4]

minmonth = reportmonth
minday = reportday

maxdt = datetime.datetime(year=year, month=int(maxmonth), day=int(maxday))
mindt = datetime.datetime(year=year, month=int(minmonth), day=int(minday))

td = maxdt - mindt
nights = td.days

ws.cell(row=2,column=2).value = reiwayear
ws.cell(row=2,column=4).value = maxmonth
ws.cell(row=2,column=6).value = maxday

ws.cell(row=7,column=14).value = reiwayear
ws.cell(row=7,column=16).value = minmonth
ws.cell(row=7,column=18).value = minday
ws.cell(row=8,column=14).value = reiwayear
ws.cell(row=8,column=16).value = maxmonth
ws.cell(row=8,column=18).value = maxday

ws.cell(row=13,column=4).value = str(reiwayear) + '年' + minmonth + '月' + minday + '日'
ws.cell(row=13,column=12).value = str(reiwayear) + '年' + maxmonth + '月' + maxday + '日'

sql = iris.sql.prepare("Select reportmonth, reportday, amount, description, paymentto from ync.expense where (reportmonth = ? and reportday >= ? and accounts = '旅費交通費') order by reportmonth, reportday")
itemline = sql.execute(reportmonth,reportday).dataframe()
	
linepos = 16

hotelamount = 0
hotelname = ''
	  
for index,row in itemline.iterrows():
	rowline = list(row)
	month = rowline[0]
	day = rowline[1]
	amount = rowline[2]
	description = rowline[3]
	paymentto = rowline[4]

	sql = iris.sql.prepare("select accounts, amount from ync.expenseitem where description = ?")
	rs = sql.execute(description)
	for index,row in enumerate(rs):
		accounts = row[0]
		if (accounts is None or accounts == ''): accounts = '旅費交通費'
		if (amount == 0 or amount is None): amount = row[1]
		
	if description == '雲雀丘花屋敷>大阪梅田': continue
	if description == '大阪梅田>雲雀丘花屋敷': continue
	
	if description == 'HOTEL':
		hotelamount = hotelamount + amount
		hotelname = hotelname + paymentto + ' ' 
		continue
	if paymentto == 'JAL':
		description = '飛行機　(' + description + ')'
	elif paymentto == 'PEACH' :
		description = '飛行機　(' + description + ')'
	elif paymentto == 'SKYMARK' :
		description = '飛行機　(' + description + ')'
	elif paymentto == 'ANA' :
		description = '飛行機　(' + description + ')'
	else:
		description = '電車　(' + description + ')'	
	linepos = linepos + 1
	if linepos > 26:
		print('# of lineitems exceeded the max limit') 
		continue
	ws.cell(row=linepos,column=1).value = str(month) + '/' + str(day)
	ws.cell(row=linepos,column=6).value = str(month) + '/' + str(day)
	ws.cell(row=linepos,column=10).value = description

	ws.cell(row=linepos,column=16).value = amount

sql = None
ws.cell(row=30,column=8).value = hotelamount
ws.cell(row=30,column=2).value = hotelname
ws.cell(row=28,column=5).value = nights
wb.save(outputFilename)
wb.close()
